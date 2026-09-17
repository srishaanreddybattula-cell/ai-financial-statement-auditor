from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def _safe_number(value):
    if value is None:
        return "N/A"
    try:
        return float(value)
    except (TypeError, ValueError):
        return "N/A"


def build_excel_report(result, company_name, ticker, cik):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"

    header_fill = PatternFill("solid", fgColor="E9EEF5")
    header_font = Font(bold=True)
    title_font = Font(bold=True, size=16)

    summary["A1"] = "Financial Reporting Risk Screening Report"
    summary["A1"].font = title_font
    summary.merge_cells("A1:D1")

    summary_rows = [
        ("Company", company_name or "N/A"),
        ("Ticker", ticker or "N/A"),
        ("CIK", cik or "N/A"),
        ("Latest annual period", result.get("latest_year", "N/A")),
        ("Screening score", _safe_number(result.get("risk_score"))),
        ("Prototype risk category", result.get("risk_category", "N/A")),
    ]

    for row_number, (label, value) in enumerate(summary_rows, start=3):
        summary.cell(row=row_number, column=1, value=label).font = header_font
        summary.cell(row=row_number, column=2, value=value)

    filing = result.get("filing") or {}
    summary["A10"] = "Source filing"
    summary["A10"].font = header_font
    source_rows = [
        ("Form", filing.get("form", "N/A")),
        ("Filed", filing.get("filing_date", "N/A")),
        ("Report date", filing.get("report_date", "N/A")),
        ("Primary document", filing.get("primary_document", "N/A")),
        ("SEC accession number", filing.get("accession_number", "N/A")),
        ("SEC source filing", filing.get("sec_url", "N/A")),
    ]
    for row_number, (label, value) in enumerate(source_rows, start=11):
        summary.cell(row=row_number, column=1, value=label).font = header_font
        cell = summary.cell(row=row_number, column=2, value=value)
        if label == "SEC source filing" and filing.get("sec_url"):
            cell.hyperlink = filing["sec_url"]
            cell.style = "Hyperlink"

    summary["A18"] = "Interpretation"
    summary["A18"].font = header_font
    summary["A19"] = (
        "This workbook is a prototype screening model for potential financial reporting risk indicators. "
        "A flagged indicator does not establish an accounting error, fraud, or material misstatement. "
        "Unavailable data is shown as N/A and excluded from the score where applicable."
    )
    summary.merge_cells("A19:D20")
    summary["A19"].alignment = Alignment(wrap_text=True, vertical="top")

    _add_risk_dimensions_sheet(workbook, result)
    _add_findings_sheet(workbook, result)
    _add_data_quality_sheet(workbook, result)
    _add_financial_data_sheet(workbook, result)
    _add_goodwill_sheet(workbook, result)
    _add_peer_sheet(workbook, result)
    _add_policies_sheet(workbook, result)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for column_cells in sheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_length = max((len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells), default=0)
            sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 60)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _style_header(sheet):
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="E9EEF5")
        cell.font = Font(bold=True)


def _add_risk_dimensions_sheet(workbook, result):
    sheet = workbook.create_sheet("Risk Dimensions")
    sheet.append(["Dimension", "Risk level"])
    _style_header(sheet)
    for name, value in result.get("risk_dimensions", {}).items():
        sheet.append([name.replace("_", " ").title(), _safe_number(value)])

    sheet.append([])
    sheet.append(["Score breakdown"])
    sheet.append(["Dimension", "Risk level", "Weight", "Contribution", "Data available"])
    for cell in sheet[sheet.max_row]:
        cell.fill = PatternFill("solid", fgColor="E9EEF5")
        cell.font = Font(bold=True)

    for item in result.get("score_breakdown", []):
        sheet.append([
            item["dimension"].replace("_", " ").title(),
            _safe_number(item.get("risk_level")),
            item.get("weight", "N/A"),
            _safe_number(item.get("contribution")),
            "Yes" if item.get("available") else "No",
        ])


def _add_findings_sheet(workbook, result):
    sheet = workbook.create_sheet("Key Findings")
    sheet.append(["Priority", "Finding", "Evidence", "Why it matters"])
    _style_header(sheet)
    for finding in result.get("findings", []):
        sheet.append([
            finding.get("priority", "N/A"),
            finding.get("title", "N/A"),
            finding.get("evidence", "N/A"),
            finding.get("why_it_matters", "N/A"),
        ])


def _add_data_quality_sheet(workbook, result):
    sheet = workbook.create_sheet("Data Quality")
    sheet.append(["Metric", "Value"])
    _style_header(sheet)
    quality = result.get("data_quality") or {}
    sheet.append(["Coverage percent", _safe_number(quality.get("coverage_percent"))])
    sheet.append(["Status", quality.get("status", "N/A")])
    sheet.append(["Missing core metrics", ", ".join(quality.get("missing", [])) or "None"])
    sheet.append(["Insufficient-history metrics", ", ".join(quality.get("insufficient_history", [])) or "None"])


def _add_financial_data_sheet(workbook, result):
    sheet = workbook.create_sheet("Financial Data")
    sheet.append(["Metric", "Fiscal year", "Value", "Form", "Filed", "Accession number", "Period start", "Period end", "Frame"])
    _style_header(sheet)
    for metric, values in result.get("financial_data", {}).items():
        if not isinstance(values, list):
            continue
        for item in values:
            sheet.append([
                metric,
                item.get("year"),
                _safe_number(item.get("value")),
                item.get("form") or "N/A",
                item.get("filed") or "N/A",
                item.get("accn") or item.get("accession_number") or "N/A",
                item.get("start") or "N/A",
                item.get("end") or "N/A",
                item.get("frame") or "N/A",
            ])


def _add_goodwill_sheet(workbook, result):
    sheet = workbook.create_sheet("Goodwill")
    sheet.append(["Metric", "Value"])
    _style_header(sheet)
    goodwill = result.get("goodwill_analysis") or {}
    sheet.append(["Goodwill / total assets", _safe_number(goodwill.get("goodwill_to_assets"))])
    sheet.append(["Year-over-year goodwill change", _safe_number(goodwill.get("goodwill_change"))])
    sheet.append(["Screening status", "Review suggested" if goodwill.get("flag") else "No threshold triggered"])
    sheet.append(["Message", goodwill.get("message", "N/A")])


def _add_peer_sheet(workbook, result):
    sheet = workbook.create_sheet("Peer Comparison")
    sheet.append(["Metric", "Company", "Peer median", "Risk direction", "Robust z-score", "Risk level"])
    _style_header(sheet)
    peer = result.get("peer_comparison") or {}
    for metric, values in peer.get("metrics", {}).items():
        sheet.append([
            metric.replace("_", " ").title(),
            _safe_number(values.get("company_value")),
            _safe_number(values.get("peer_median")),
            values.get("risk_direction") or "N/A",
            _safe_number(values.get("z_score")),
            _safe_number(values.get("risk_level")),
        ])


def _add_policies_sheet(workbook, result):
    sheet = workbook.create_sheet("Accounting Policies")
    sheet.append(["Policy", "Matches"])
    _style_header(sheet)
    policy = result.get("policy_analysis") or {}
    for name, matches in policy.get("policy_matches", {}).items():
        sheet.append([name.replace("_", " ").title(), ", ".join(matches)])
